import openpyxl
from pathlib import Path
import os
from dotenv import load_dotenv

load_dotenv()

BASE_PATH_STR = os.getenv("GS_BASE_PATH")
if not BASE_PATH_STR:
    raise ValueError("Variable GS_BASE_PATH no encontrada en .env")

BASE_PATH = Path(BASE_PATH_STR)

RUTA_CENTRALES = BASE_PATH / "01-Bulk Merge_Centrales_Ids.xlsx"
RUTA_PORTS = BASE_PATH / "01-Bulk export of ports.xlsx"

# ------------------------------------------------------------
# Funciones auxiliares
# ------------------------------------------------------------
def normalizar_id(valor):
    if valor is None or valor == "":
        return ""
    if isinstance(valor, (int, float)):
        try:
            return str(int(valor))
        except ValueError:
            return str(valor).strip()
    return str(valor).strip()

def extraer_indices(worksheet):
    """Devuelve diccionario {nombre_columna_lower: índice (0-based)}"""
    indices = {}
    for idx, cell in enumerate(worksheet[1], start=0):
        if cell.value:
            val = str(cell.value).strip().lower()
            indices[val] = idx
    return indices

# ------------------------------------------------------------
# 1. Cargar archivo de centrales (hoja "IS") como diccionario EC -> set(anillos)
# ------------------------------------------------------------
def cargar_centrales(ruta):
    """
    Lee la hoja "IS" y devuelve:
    - flat_pairs: set de tuplas (EC, ANILLO) para validación rápida
    - ec_to_anillos: dict {EC: set(anillos)} para mensajes detallados
    """
    flat_pairs = set()
    ec_to_anillos = {}
    wb = None
    try:
        wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
        
        if "IS" not in wb.sheetnames:
            print(f"Hoja 'IS' no encontrada. Hojas disponibles: {wb.sheetnames}")
            return flat_pairs, ec_to_anillos
        
        ws = wb["IS"]
        idx = extraer_indices(ws)
        
        col_ec = idx.get('ec')
        col_anillo = idx.get('anillo')
        
        if col_ec is None or col_anillo is None:
            print("No se encontraron columnas 'ec' o 'anillo' en la hoja IS.")
            return flat_pairs, ec_to_anillos

        for row in ws.iter_rows(min_row=2, values_only=True):
            ec = normalizar_id(row[col_ec]) if col_ec < len(row) else ""
            anillo = str(row[col_anillo]).strip() if col_anillo < len(row) and row[col_anillo] else ""
            
            if ec and anillo:
                flat_pairs.add((ec, anillo))
                if ec not in ec_to_anillos:
                    ec_to_anillos[ec] = set()
                ec_to_anillos[ec].add(anillo)

        print(f"Centrales cargadas desde hoja 'IS': {len(flat_pairs)} pares (EC, ANILLO)")
        print(f"EC únicos encontrados: {len(ec_to_anillos)}")
    except Exception as e:
        print(f"Error al cargar centrales: {e}")
    finally:
        if wb:
            wb.close()
    return flat_pairs, ec_to_anillos

# ------------------------------------------------------------
# 2. Validar puertos contra centrales (con descripción detallada)
# ------------------------------------------------------------
def validar_puertos_contra_centrales(ruta_ports, flat_pairs, ec_to_anillos):
    """
    Itera sobre los puertos y verifica que el par (Equipo Central, Proyecto De Red)
    exista en el conjunto de centrales.
    Retorna lista de errores: (id_port, shelf_desc, proyecto, equipo, mensaje)
    """
    errores = []
    wb = None
    try:
        wb = openpyxl.load_workbook(ruta_ports, data_only=True, read_only=True)
        ws = wb["Port"] if "Port" in wb.sheetnames else wb.worksheets[0]
        idx = extraer_indices(ws)

        # --- Buscar columnas ---
        col_id = None
        col_proyecto = None
        col_equipo = None
        col_shelf_desc = None

        for col_name, pos in idx.items():
            # ID del puerto (columna "Id")
            if col_name == 'id':
                col_id = pos
            # Proyecto de Red
            if 'proyecto de red' in col_name:
                col_proyecto = pos
            # Equipo Central
            if 'equipo central' in col_name:
                col_equipo = pos
            # Description Shelf: puede ser "description (3)" o "descripcion shelf"
            if 'description (3)' in col_name or 'descripcion shelf' in col_name:
                col_shelf_desc = pos

        # Si no encuentra "description (3)", buscar cualquier columna que contenga "shelf" y "desc"
        if col_shelf_desc is None:
            for col_name, pos in idx.items():
                if 'shelf' in col_name and ('desc' in col_name or 'description' in col_name):
                    col_shelf_desc = pos
                    break

        if col_id is None or col_proyecto is None or col_equipo is None:
            print("No se encontraron columnas 'Id', 'Proyecto De Red' y/o 'Equipo Central'.")
            return errores

        # --- Recorrer puertos ---
        for row in ws.iter_rows(min_row=2, values_only=True):
            id_port = normalizar_id(row[col_id]) if col_id < len(row) else ""
            proyecto = str(row[col_proyecto]).strip() if col_proyecto < len(row) and row[col_proyecto] else ""
            equipo = str(row[col_equipo]).strip() if col_equipo < len(row) and row[col_equipo] else ""
            shelf_desc = str(row[col_shelf_desc]).strip() if col_shelf_desc is not None and col_shelf_desc < len(row) and row[col_shelf_desc] else ""

            if id_port and proyecto and equipo:
                if (equipo, proyecto) not in flat_pairs:
                    # Verificar si el EC existe en el bulk
                    if equipo not in ec_to_anillos:
                        mensaje = (f"EC '{equipo}' NO EXISTE en el archivo de centrales (hoja IS). "
                                   f"Verifique el nombre del Equipo Central en el puerto.")
                    else:
                        anillos_esperados = sorted(ec_to_anillos[equipo])
                        # Limitar a 5 para no hacer el mensaje demasiado largo
                        anillos_str = ", ".join(anillos_esperados[:5])
                        if len(anillos_esperados) > 5:
                            anillos_str += f", ... (total {len(anillos_esperados)})"
                        mensaje = (f"ANILLO INCORRECTO para el EC '{equipo}'. "
                                   f"En la hoja IS del Bulk, este EC está asociado a: [{anillos_str}]. "
                                   f"Sin embargo, en el puerto se encontró: '{proyecto}'.")
                    
                    errores.append((id_port, shelf_desc, proyecto, equipo, mensaje))

    except Exception as e:
        print(f"Error al validar puertos: {e}")
        import traceback
        traceback.print_exc()
    finally:
        if wb:
            wb.close()
    return errores

# ------------------------------------------------------------
# 3. Generar reporte A1.xlsx con columnas mejoradas
# ------------------------------------------------------------
def generar_reporte(errores, ruta_salida):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Errores_Anillo_EC"
    # Nuevas columnas: ID Puerto | Description Shelf | Proyecto De Red | Equipo Central | Error Detallado
    ws.append(["ID Puerto", "Description Shelf", "Proyecto De Red", "Equipo Central", "Error"])

    for id_port, shelf_desc, proyecto, equipo, msg in errores:
        ws.append([id_port, shelf_desc, proyecto, equipo, msg])

    # Ajustar ancho de columnas
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)  # más ancho para el error

    wb.save(ruta_salida)
    print(f"Reporte guardado en: {ruta_salida}")

# ------------------------------------------------------------
# 4. Ejecución principal
# ------------------------------------------------------------
if __name__ == "__main__":
    print("=" * 65)
    print("  VALIDACIÓN DE ANILLOS vs CENTRALES (TEST - DETALLADO)")
    print("=" * 65)

    if not RUTA_CENTRALES.exists():
        print(f"ERROR: No se encuentra {RUTA_CENTRALES}")
        exit(1)
    if not RUTA_PORTS.exists():
        print(f"ERROR: No se encuentra {RUTA_PORTS}")
        exit(1)

    # Cargar datos de centrales
    flat_pairs, ec_to_anillos = cargar_centrales(RUTA_CENTRALES)
    if not flat_pairs:
        print("No se cargaron pares de centrales. Revise el archivo.")
        exit(1)

    # Validar puertos
    errores = validar_puertos_contra_centrales(RUTA_PORTS, flat_pairs, ec_to_anillos)
    print(f"\nPuertos procesados. Errores encontrados: {len(errores)}")

    # Generar reporte
    ruta_reporte = BASE_PATH / "A1.xlsx"
    generar_reporte(errores, ruta_reporte)

    print("=" * 65)
    print("  FIN DEL TEST")
    print("=" * 65)