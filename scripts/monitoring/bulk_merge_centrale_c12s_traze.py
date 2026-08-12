#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
Extrae conexiones EC-OP-OT desde archivos de traza siguiendo la lógica de la macro VBA.
- Agrupa por contexto (indentación).
- Ordena EC por Slot ascendente.
- Asigna OP y OT en orden de aparición.
- port_ec = slot/0/puerto
- odf_patcheo = número del Shelf OP (ej. OP..._N -> N)
- port_pacheo = puertos OP en orden ascendente separados por '+'
- cassetera = letra de BANDEJA del primer OP
- odf_troncal = identificador del OT (ej. OT..._XXX -> XXX)
"""

import re
from pathlib import Path
from typing import List, Dict, Set, Optional, Tuple
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

# ----------------------------------------------------------------------
# CONFIGURACIÓN
# ----------------------------------------------------------------------
CARPETA_PADRE = Path(
    r"\\atlas\VP_INFRAESTRUCTURA\G_Plan_Gestion_Proyectos"
    r"\RI200_A_Admon_Capacidad_P_R_Infr\2023\REPORTES_INF_CART_GEO"
    r"\DICCIONARIO_DE_DATOS\40_CABLES_CORPORATIVOS\11_Trace_Conectividad\d"
)

RUTA_EXCLUSIONES = Path(
    r"C:\A_GS1_PROYECTOS\0_Documents_gs\output\quality\03-TRAZE_MS_CONNECTION.xlsx"
)

RUTA_SALIDA = Path(r"C:\A_GS1_PROYECTOS") / "Tabla_Conexiones_EC_OP_OT_final.xlsx"

# ----------------------------------------------------------------------
# FUNCIONES AUXILIARES
# ----------------------------------------------------------------------
def cargar_exclusiones(ruta: Path) -> Set[str]:
    if not ruta.exists():
        return set()
    try:
        df = pd.read_excel(ruta, sheet_name="OM", engine="openpyxl", dtype=str, header=None)
        exclusions = set()
        for col in df.columns[:2]:
            for val in df[col].dropna():
                s = str(val).strip()
                if s:
                    exclusions.add(s)
        return exclusions
    except Exception:
        return set()

def extraer_metadatos(linea: str) -> Dict[str, str]:
    datos = {}
    m = re.search(r'Port \((.*?)\)', linea)
    datos['port'] = m.group(1).strip() if m else ''
    m = re.search(r'Shelf \((.*?)\)', linea)
    datos['shelf'] = m.group(1).strip() if m else ''
    m = re.search(r'Bay \((.*?)\)', linea)
    datos['bay'] = m.group(1).strip() if m else ''
    m = re.search(r'Card \((.*?)\)', linea)
    datos['card'] = m.group(1).strip() if m else ''
    m = re.search(r'Rack \((.*?)\)', linea)
    datos['rack'] = m.group(1).strip() if m else ''
    m = re.search(r'Room \((.*?)\)', linea)
    datos['room'] = m.group(1).strip() if m else ''
    return datos

def extraer_anillo_cable_hilo(lineas: List[str]) -> Tuple[str, str, str]:
    anillo = ''
    cable = ''
    hilo = ''
    for linea in lineas:
        if 'Fibre (' in linea:
            m = re.search(r'Fibre \((.*?)\)', linea)
            if m:
                contenido = m.group(1).strip()
                if ':' in contenido:
                    parts = contenido.split(':', 1)
                    anillo = parts[0].strip().upper()
                    hilo_num = parts[1].strip()
                    hilo = re.sub(r'\D', '', hilo_num)
            break
    for linea in lineas:
        if 'Cable de FO (' in linea:
            m = re.search(r"'([^']*)'", linea)
            if m:
                cable = m.group(1).strip()
                break
    return anillo, cable, hilo

def clasificar_puerto(port: str, shelf: str) -> str:
    if '-' in port:
        return 'EC'
    if shelf.upper().startswith('OT'):
        return 'OT'
    if shelf.upper().startswith('OP'):
        return 'OP'
    return 'OTRO'

def extraer_slot(card: str) -> str:
    m = re.search(r'Slot(\d+)', card)
    return m.group(1) if m else '0'

def extraer_numero_port(port: str) -> str:
    return re.sub(r'\D', '', port)

def extraer_odf_patcheo(shelf: str) -> str:
    """Extrae el número después del último '_' del Shelf OP (ej. OP..._8 -> 8)"""
    if '_' in shelf:
        parts = shelf.split('_')
        return parts[-1] if parts[-1].isdigit() else ''
    return ''

def extraer_cassetera(card: str) -> str:
    m = re.search(r'BANDEJA ([A-Z])', card)
    return m.group(1) if m else ''

def extraer_odf_troncal(shelf: str) -> str:
    shelf_clean = shelf.split()[0] if shelf else ''
    m = re.search(r'OT\d+_(.+)', shelf_clean)
    return m.group(1) if m else shelf_clean

# ----------------------------------------------------------------------
# PROCESAR ARCHIVO
# ----------------------------------------------------------------------
def procesar_archivo(ruta: Path) -> List[Dict]:
    with open(ruta, 'r', encoding='utf-8', errors='ignore') as f:
        lineas = [l.rstrip('\n') for l in f]
    if not lineas:
        return []

    anillo, cable, hilo = extraer_anillo_cable_hilo(lineas)

    # Extraer puertos EC, OP, OT con indentación
    puertos = []
    for linea in lineas:
        if 'Port (' not in linea:
            continue
        indent = len(linea) - len(linea.lstrip())
        meta = extraer_metadatos(linea)
        port = meta['port']
        shelf = meta['shelf']
        tipo = clasificar_puerto(port, shelf)
        if tipo == 'OTRO':
            continue
        puertos.append({
            'tipo': tipo,
            'port': port,
            'shelf': shelf,
            'bay': meta['bay'],
            'card': meta['card'],
            'rack': meta['rack'],
            'room': meta['room'],
            'indent': indent,
            'linea': linea.strip()
        })

    if not puertos:
        return []

    # Asignar contexto por indentación (cambio a indentación menor -> nuevo contexto)
    contexto_actual = 0
    indent_actual = puertos[0]['indent']
    for p in puertos:
        if p['indent'] < indent_actual:
            contexto_actual += 1
        indent_actual = p['indent']
        p['contexto'] = contexto_actual

    # Agrupar por contexto
    grupos = {}
    for p in puertos:
        ctx = p['contexto']
        if ctx not in grupos:
            grupos[ctx] = {'EC': [], 'OP': [], 'OT': []}
        grupos[ctx][p['tipo']].append(p)

    resultados = []
    for ctx, tipos in grupos.items():
        ecs = tipos['EC']
        ops = tipos['OP']
        ots = tipos['OT']
        if not ecs or not ops or not ots:
            continue

        # Ordenar EC por Slot ascendente (para que el EC con Slot menor vaya primero)
        for ec in ecs:
            slot = int(extraer_slot(ec['card']) or 0)
            ec['slot_num'] = slot
        ecs_sorted = sorted(ecs, key=lambda x: x['slot_num'])

        # Distribuir OP y OT en el mismo orden que los ECs (reparto equitativo)
        ops_por_ec = len(ops) // len(ecs_sorted)
        ots_por_ec = len(ots) // len(ecs_sorted)

        for i, ec in enumerate(ecs_sorted):
            inicio_op = i * ops_por_ec
            fin_op = inicio_op + ops_por_ec if i < len(ecs_sorted)-1 else len(ops)
            ops_actuales = ops[inicio_op:fin_op]
            if not ops_actuales:
                continue

            inicio_ot = i * ots_por_ec
            ot = ots[inicio_ot] if inicio_ot < len(ots) else None
            if not ot:
                continue

            # port_ec: slot/0/puerto
            slot = extraer_slot(ec['card'])
            port_num = extraer_numero_port(ec['port'])
            port_ec = f"{slot}/0/{port_num}" if slot and port_num else ec['port']

            # port_pacheo: puertos OP en orden ascendente (numérico)
            ports_op = [op['port'] for op in ops_actuales]
            ports_op_sorted = sorted(ports_op, key=lambda x: int(x) if x.isdigit() else 0)
            port_pacheo = '+'.join(ports_op_sorted) if ports_op_sorted else ''

            # odf_patcheo: número del Shelf del primer OP (OP..._N)
            odf_patcheo = extraer_odf_patcheo(ops_actuales[0]['shelf'])

            # cassetera: letra de BANDEJA del primer OP
            cassetera = extraer_cassetera(ops_actuales[0]['card'])

            # odf_troncal: identificador del OT
            odf_troncal = extraer_odf_troncal(ot['shelf'])

            fila = {
                'ec': ec['shelf'],
                'port_ec': port_ec,
                'odf_patcheo': odf_patcheo,
                'cassetera': cassetera,
                'port_pacheo': port_pacheo,
                'fibra': cable,
                'odf_troncal': odf_troncal,
                'hilo': hilo,
                'anillo': anillo,
            }
            resultados.append(fila)

    return resultados

# ----------------------------------------------------------------------
# RECORRIDO Y GENERACIÓN DE DATAFRAME
# ----------------------------------------------------------------------
def generar_tabla(carpeta_padre: Path, exclusiones: Set[str]) -> pd.DataFrame:
    print(f"Buscando archivos en: {carpeta_padre}")
    if not carpeta_padre.exists():
        print("❌ La carpeta NO EXISTE.")
        return pd.DataFrame(columns=['ec', 'port_ec', 'odf_patcheo', 'cassetera', 'port_pacheo', 'fibra', 'odf_troncal', 'hilo', 'anillo'])

    archivos = set()
    for ext in ['*.txt', '*.TXT']:
        for p in carpeta_padre.rglob(ext):
            archivos.add(p)

    archivos_filtrados = [a for a in archivos if a.stem not in exclusiones]
    print(f"Archivos a procesar: {len(archivos_filtrados)}")

    datos = []
    for arch in archivos_filtrados:
        filas = procesar_archivo(arch)
        if filas:
            datos.extend(filas)

    if not datos:
        print("No se encontraron conexiones completas.")
        return pd.DataFrame(columns=['ec', 'port_ec', 'odf_patcheo', 'cassetera', 'port_pacheo', 'fibra', 'odf_troncal', 'hilo', 'anillo'])

    df = pd.DataFrame(datos)
    df = df.drop_duplicates().sort_values(['anillo', 'fibra', 'hilo', 'ec']).reset_index(drop=True)
    print(f"✅ Se generaron {len(df)} filas.")
    return df

# ----------------------------------------------------------------------
# GUARDAR EN EXCEL
# ----------------------------------------------------------------------
def guardar_excel(df: pd.DataFrame, ruta_salida: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Conexiones"

    if df.empty:
        ws.cell(row=1, column=1, value="No se encontraron conexiones.")
        wb.save(ruta_salida)
        print(f"Archivo guardado (vacío) en {ruta_salida}")
        return

    headers = ['ec', 'port_ec', 'odf_patcheo', 'cassetera', 'port_pacheo', 'fibra', 'odf_troncal', 'hilo', 'anillo']
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)

    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        ws.cell(row=row_idx, column=1, value=row[0])
        ws.cell(row=row_idx, column=2, value=row[1])
        ws.cell(row=row_idx, column=3, value=row[2])
        ws.cell(row=row_idx, column=4, value=row[3])
        ws.cell(row=row_idx, column=5, value=row[4])
        ws.cell(row=row_idx, column=6, value=row[5])
        ws.cell(row=row_idx, column=7, value=row[6])
        ws.cell(row=row_idx, column=8, value=row[7])
        ws.cell(row=row_idx, column=9, value=row[8])

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 30)

    wb.save(ruta_salida)
    print(f" Tabla guardada en {ruta_salida}")

# ----------------------------------------------------------------------
# MAIN
# ----------------------------------------------------------------------
def main():
    print("=== EXTRACCIÓN DE CONEXIONES EC-OP-OT (macro exacta) ===")
    exclusiones = cargar_exclusiones(RUTA_EXCLUSIONES)
    print(f"Exclusiones: {len(exclusiones)} elementos")

    df = generar_tabla(CARPETA_PADRE, exclusiones)
    guardar_excel(df, RUTA_SALIDA)

if __name__ == "__main__":
    main()